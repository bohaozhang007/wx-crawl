from pathlib import Path
import math
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from PIL import Image, ImageDraw, ImageFont

OUT=Path('/root/workspace/wx-crawl/slides');OUT.mkdir(exist_ok=True)
r=Presentation();r.slide_width=Inches(13.333);r.slide_height=Inches(7.5)
BG='F4F6F8';INK='253442';MUTED='70808F';ACC='607F98';LINE='CDD6DE';WHITE='FFFFFF';PALE='E6EDF3'
FONT='Microsoft YaHei'; slides=[]
def rgb(c):return RGBColor.from_string(c)
def slide(n,section,title,sub):
 s=r.slides.add_slide(r.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(BG)
 slides.append([])
 text(s,.7,.34,9,.25,'WX-CRAWL   /   '+section,10,MUTED)
 text(s,.7,.96,12,.7,title,30,INK,True)
 text(s,.73,1.8,11.8,.52,sub,14,MUTED)
 line(s,.72,6.96,12.6,6.96,False)
 text(s,.73,7.08,10,.2,'微信公众号科研机会  ·  工具调用与信息筛选',9,MUTED)
 text(s,12,7.05,.6,.24,f'{n:02d} / 05',9,MUTED)
 return s
def box(s,x,y,w,h,fill=WHITE,border=LINE):
 q=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));q.fill.solid();q.fill.fore_color.rgb=rgb(fill)
 q.line.color.rgb=rgb(border);q.line.width=Pt(.7)
 slides[-1].append(('box',x,y,w,h,fill,border));return q
def text(s,x,y,w,h,t,size=16,color=INK,bold=False):
 q=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=q.text_frame;tf.word_wrap=True;tf.margin_left=tf.margin_right=0;tf.margin_top=Pt(1);tf.margin_bottom=0
 for i,st in enumerate(t.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=st;p.font.name=FONT;p.font.size=Pt(size);p.font.color.rgb=rgb(color);p.font.bold=bold;p.space_after=Pt(8)
  for run in p.runs:
   ea=OxmlElement('a:ea');ea.set('typeface',FONT);run._r.get_or_add_rPr().append(ea)
 slides[-1].append(('text',x,y,w,h,t,size,color,bold));return q
def line(s,x1,y1,x2,y2,arrow=True):
 q=s.shapes.add_connector(1,Inches(x1),Inches(y1),Inches(x2),Inches(y2));q.line.color.rgb=rgb(ACC);q.line.width=Pt(1.1)
 if arrow:
  end=OxmlElement('a:tailEnd');end.set('type','triangle');q._element.spPr.get_or_add_ln().append(end)
 slides[-1].append(('line',x1,y1,x2,y2,arrow))
def card(s,x,y,w,h,num,title,body):
 box(s,x,y,w,h);text(s,x+.18,y+.18,w-.36,.25,num,10,ACC);text(s,x+.18,y+.65,w-.36,.4,title,18,INK,True);text(s,x+.18,y+1.25,w-.36,h-1.3,body,13,MUTED)

s=slide(1,'PIPELINE','从公众号文章，到可行动的科研机会','列表发现 → 正文归档 → 规则判断 → 机会入库 → 消息提醒')
items=[('01','历史列表','TikHub V2 HTTP\n分页获取 · 增量去重'),('02','正文下载','两级开源工具\nTikHub 最后兜底'),('03','打标与筛选','地域 · 截止 · 方向\nKEEP / DROP / REVIEW'),('04','入库与同步','SQLite 保存标签依据\n同步钉钉 AI 表'),('05','机会提醒','高重要度提醒\n申报截止提醒')]
for i,(num,t,b) in enumerate(items):
 x=.73+i*2.43;card(s,x,2.7,2.15,2.7,num,t,b)
 if i<4:line(s,x+2.16,3.98,x+2.4,3.98)
text(s,.76,5.93,11.8,.4,'关注范围   全国性国家级项目 / 北京市 / 浙江省（含下辖市区县）',16,ACC)

s=slide(2,'COLLECTION','历史列表付费获取，正文按序兜底','当前正式历史后端：TikHub；正文成功通过校验后，立即停止后续下载尝试。')
box(s,.73,2.58,11.87,.83,PALE,PALE)
text(s,.95,2.83,11.45,.3,'run.sh   →   公众号身份核验   →   TikHub 历史列表   →   URL 去重（归档 + SQLite）',16)
for i,(title,desc) in enumerate([('wechat-mp-tools','第一优先 · 开源下载'),('we-mp-rss','第二优先 · 开源下载'),('TikHub 正文接口','最后兜底 · 付费调用')]):
 x=.73+i*4.05;box(s,x,3.99,3.76,1.35);text(s,x+.2,4.19,3.36,.4,title,20,INK,True);text(s,x+.2,4.81,3.36,.28,desc,13,MUTED)
 if i<2:line(s,x+3.77,4.63,x+4.02,4.63)
text(s,.78,3.59,11,.27,'向右切换条件：下载失败或正文校验不通过',12,ACC)
text(s,.78,5.76,11.8,.7,'通过校验 → 归档 HTML、文本、元数据；全部失败 → 记录失败，保留待处理。\n图片型正文可通过下载校验，但不代表图片文字已经识别。',14,MUTED)

s=slide(3,'SELECTION','先确认有效范围，再判断科研价值','模型生成结构化标签；程序在筛选与发送时再次执行门禁。')
rows=[('01  版本与证据','当前规则 + 有效结构 + 原文证据','不合格 → 重试 / 待处理'),('02  截止有效性','未截止，或截止时间不明确','明确已截止 → 排除'),('03  地域范围','全国性国家级 / 北京 / 浙江','其他地区 → 排除；不明 → REVIEW'),('04  机会与方向','项目或指南申报 + 目标研究方向','KEEP 入选；DROP 排除；REVIEW 复核')]
for i,(a,b,c) in enumerate(rows):
 y=2.57+i*.79;box(s,.73,y,11.87,.67,WHITE,WHITE);text(s,.93,y+.17,2.3,.3,a,15,INK,True);text(s,3.28,y+.17,4.6,.3,b,14);text(s,8.02,y+.17,4.33,.3,c,12,MUTED)
text(s,.82,6.1,11.7,.42,'例：国家自然科学基金区域联合基金（江西）属于其他地区专项，不能因“国家”字样放行。',13,ACC)

s=slide(4,'DELIVERY','入库与提醒分支执行，发送前再校验','符合范围的 KEEP 机会进入下游；高重要度提醒不以数据库同步完成为前提。')
box(s,.75,3.5,2,1.25,PALE,PALE);text(s,.95,3.81,1.6,.6,'有效 KEEP\n科研机会',20,INK,True)
line(s,2.75,4.12,3.1,4.12,False);line(s,3.1,2.94,3.1,5.6,False)
for y,t,b in [(2.5,'高重要度提醒','high → 当前门禁复核 → 幂等去重 → 钉钉提醒'),(3.9,'入库与同步','筛选报告 → 入库前复核 → SQLite → 钉钉 AI 表'),(5.3,'截止提醒','库中未来截止时间 → 排期 → 发送前复核 → 钉钉提醒')]:
 line(s,3.1,y+.43,3.52,y+.43);box(s,3.55,y,9,.93);text(s,3.78,y+.14,2.2,.34,t,17,INK,True);text(s,6.08,y+.24,6.15,.38,b,12,MUTED)

s=slide(5,'CONTROL','规则升级必须可追溯','当前已实施契约：标签格式 schema 2 / 决策树 tree 1.2 / 领域配置 profile 1.3')
for i,(n,t,b) in enumerate([('01','读取当前契约','以规则文件为准\n不依赖历史完成标记'),('02','验证历史结果','旧版本或缺地域依据\n必须实际重新判定'),('03','保留判断依据','地域结论与原文证据\n随完整标签一同保存')]):
 card(s,.73+i*4.05,2.66,3.76,2.7,n,t,b)
box(s,.73,5.78,11.86,.65,PALE,PALE)
text(s,.94,5.96,11.4,.3,'截止不明可继续筛选；地域不明须复核。数量一致，不等于内容与规则版本合规。',15,ACC)

r.core_properties.title='微信公众号工具调用与信息筛选流程'
r.core_properties.subject='wx-crawl 工作流｜冷色极简版'
r.core_properties.author=''
r.save(OUT/'微信公众号工具调用与信息筛选流程.pptx')
# Validate editable structure and slide bounds.
check=Presentation(OUT/'微信公众号工具调用与信息筛选流程.pptx')
assert len(check.slides)==5
for s in check.slides:
 for q in s.shapes:
  assert q.left>=0 and q.top>=0 and q.left+q.width<=r.slide_width+100 and q.top+q.height<=r.slide_height+100
fontpath=Path('/tmp/NotoSansCJKsc-Regular.otf')
if fontpath.exists() and fontpath.stat().st_size>100000:
 imgs=[];scale=120
 for idx,ops in enumerate(slides):
  im=Image.new('RGB',(1600,900),'#'+BG);d=ImageDraw.Draw(im)
  for op in ops:
   if op[0]=='box':
    _,x,y,w,h,f,b=op;d.rectangle((x*scale,y*scale,(x+w)*scale,(y+h)*scale),fill='#'+f,outline='#'+b,width=1)
   elif op[0]=='line':
    _,x,y,xx,yy,arrow=op;d.line((x*scale,y*scale,xx*scale,yy*scale),fill='#'+ACC,width=2)
    if arrow:d.polygon([(xx*scale,yy*scale),(xx*scale-8,yy*scale-4),(xx*scale-8,yy*scale+4)],fill='#'+ACC)
   else:
    _,x,y,w,h,t,size,col,bold=op;ft=ImageFont.truetype(str(fontpath),int(size/72*scale));y0=y*scale
    for ln in t.split('\n'):
     d.text((x*scale,y0),ln,font=ft,fill='#'+col);y0+=size/72*scale+8/72*scale
  im.save(OUT/f'slide-{idx+1}.png');imgs.append(im)
 contact=Image.new('RGB',(1600,3*450),'#'+LINE)
 for i,im in enumerate(imgs):contact.paste(im.resize((800,450)),((i%2)*800,(i//2)*450))
 contact.save(OUT/'preview.png')
print(OUT/'微信公众号工具调用与信息筛选流程.pptx')
