from pathlib import Path
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from PIL import Image,ImageDraw,ImageFont
O=Path('/root/workspace/wx-crawl/slides/report');R=Presentation();R.slide_width=Inches(13.333);R.slide_height=Inches(7.5)
s=R.slides.add_slide(R.slide_layouts[6]);BG='F5F7F9';INK='263746';MUT='6C7F8E';BLUE='537C9B';LINE='D3DFE7';PALE='E8EFF4';s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(BG)
im=Image.new('RGB',(1600,900),'#'+BG);d=ImageDraw.Draw(im);scale=120;fontpath='/tmp/NotoSansCJKsc-Regular.otf'
def txt(x,y,w,h,t,sz=16,col=INK,bold=False):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.word_wrap=True;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 ft=ImageFont.truetype(fontpath,round(sz*scale/72))
 for i,st in enumerate(t.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=st;p.font.name='Microsoft YaHei';p.font.size=Pt(sz);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(col);p.space_after=Pt(9)
  for run in p.runs:
   ea=OxmlElement('a:ea');ea.set('typeface','Microsoft YaHei');run._r.get_or_add_rPr().append(ea)
  d.text((x*scale,y*scale+i*(sz+9)*scale/72),st,font=ft,fill='#'+col)
def box(x,y,w,h,fill='FFFFFF',border=LINE):
 sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill);sh.line.color.rgb=RGBColor.from_string(border);sh.line.width=Pt(.6)
 d.rectangle((x*scale,y*scale,(x+w)*scale,(y+h)*scale),fill='#'+fill,outline='#'+border,width=1)
def arrow(x,y,xx):
 sh=s.shapes.add_connector(1,Inches(x),Inches(y),Inches(xx),Inches(y));sh.line.color.rgb=RGBColor.from_string(BLUE);sh.line.width=Pt(1.2);end=OxmlElement('a:tailEnd');end.set('type','triangle');sh._element.spPr.get_or_add_ln().append(end)
 d.line((x*scale,y*scale,xx*scale,y*scale),fill='#'+BLUE,width=2);d.polygon([(xx*scale,y*scale),(xx*scale-7,y*scale-4),(xx*scale-7,y*scale+4)],fill='#'+BLUE)
txt(.65,.38,12,.3,'RESEARCH OPPORTUNITY  /  AUTOMATED PIPELINE',10,MUT)
txt(.65,.94,12,.7,'微信公众号科研机会 · 自动化采集与筛选',29,INK,True)
txt(.68,1.74,12,.4,'将分散的公众号信息，转化为可查询、可跟进的科研项目与指南申报机会。',16,MUT)
box(.68,2.43,11.97,.58,PALE,PALE)
txt(.88,2.56,11.55,.3,'任务编排：Hermes 定时触发 + Python 串联处理流程',15,BLUE)
items=[('01','列表采集','TikHub','获取公众号历史列表\n增量发现新文章'),('02','正文获取','wechat-mp-tools\nwe-mp-rss','下载、归档文章正文\n失败时由 TikHub 兜底'),('03','智能打标','DeepSeek 大模型','识别类型、研究方向\n提取截止时间、重要度'),('04','规则筛选','Python 规则校验','按地域、时效筛选\n不确定项转人工复核'),('05','入库与推送','SQLite + 钉钉','结构化存储、AI 表同步\n重要机会及截止提醒')]
for i,(num,title,tool,body) in enumerate(items):
 x=.68+i*2.44;box(x,3.43,2.21,2.55)
 txt(x+.16,3.61,1.89,.25,num,10,BLUE)
 txt(x+.16,4.02,1.9,.4,title,20,INK,True)
 txt(x+.16,4.59,1.93,.63,tool,13,BLUE,True)
 txt(x+.16,5.3,1.93,.56,body,11.7,MUT)
 if i<4:arrow(x+2.24,4.57,x+2.41)
txt(.72,6.44,12,.4,'关注范围：全国性国家级项目、北京市、浙江省（含下辖市区县）',16,INK)
txt(.72,7.05,12,.22,'开源正文优先，付费接口兜底  ·  大模型识别，明确规则把关',10,MUT)
R.core_properties.title='科研机会自动化 Pipeline｜汇报版';R.core_properties.subject='公众号采集、智能打标、规则筛选、入库与推送'
p=O/'科研机会自动化Pipeline-一页汇报版.pptx';R.save(p);im.save(O/'preview.png')
check=Presentation(p);assert len(check.slides)==1
for sh in check.slides[0].shapes:assert sh.left>=0 and sh.top>=0 and sh.left+sh.width<=R.slide_width and sh.top+sh.height<=R.slide_height
print(p)
